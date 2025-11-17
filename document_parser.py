"""
Document Parser Module

Supports parsing of PDF, DOCX, and TXT files with structure preservation
and intelligent semantic chunking for compliance analysis.
"""

import os
import re
from pathlib import Path
from typing import List, Dict, Optional, Any
from datetime import datetime
from dataclasses import dataclass

try:
    import PyPDF2
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


@dataclass
class DocumentChunk:
    """Represents a chunk of text with metadata."""
    content: str
    chunk_id: int
    section: Optional[str] = None
    page_number: Optional[int] = None
    start_char: int = 0
    end_char: int = 0


@dataclass
class ParsedDocument:
    """Represents a parsed document with structure and metadata."""
    filename: str
    file_type: str
    title: Optional[str] = None
    author: Optional[str] = None
    date: Optional[datetime] = None
    full_text: str = ""
    chunks: List[DocumentChunk] = None
    metadata: Dict[str, Any] = None
    
    def __post_init__(self):
        if self.chunks is None:
            self.chunks = []
        if self.metadata is None:
            self.metadata = {}


class DocumentParser:
    """
    Parses various document formats (PDF, DOCX, TXT) with structure preservation
    and intelligent semantic chunking for compliance analysis.
    """
    
    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        semantic_chunking: bool = True
    ):
        """
        Initialize the document parser.
        
        Args:
            chunk_size: Target size for text chunks (characters)
            chunk_overlap: Overlap between chunks (characters)
            semantic_chunking: Use semantic boundaries (paragraphs, sections) when possible
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.semantic_chunking = semantic_chunking
        
        # Validate dependencies
        if not PDF_AVAILABLE:
            print("Warning: PyPDF2 not available. PDF parsing will not work.")
        if not DOCX_AVAILABLE:
            print("Warning: python-docx not available. DOCX parsing will not work.")
    
    def parse_file(self, file_path: str) -> ParsedDocument:
        """
        Parse a document file and extract content with structure.
        
        Args:
            file_path: Path to the document file
        
        Returns:
            ParsedDocument object with extracted content and metadata
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = file_path.suffix.lower()
        
        if file_ext == '.pdf':
            return self._parse_pdf(file_path)
        elif file_ext in ['.docx', '.doc']:
            return self._parse_docx(file_path)
        elif file_ext == '.txt':
            return self._parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_ext}")
    
    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """Parse PDF file."""
        if not PDF_AVAILABLE:
            raise ImportError("PyPDF2 is required for PDF parsing. Install with: pip install PyPDF2")
        
        doc = ParsedDocument(
            filename=file_path.name,
            file_type="PDF"
        )
        
        full_text_parts = []
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                # Extract metadata
                if pdf_reader.metadata:
                    doc.title = pdf_reader.metadata.get('/Title', '')
                    doc.author = pdf_reader.metadata.get('/Author', '')
                    
                    # Try to parse date
                    date_str = pdf_reader.metadata.get('/CreationDate', '')
                    if date_str:
                        doc.date = self._parse_pdf_date(date_str)
                
                # Extract text from each page
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    text = page.extract_text()
                    if text.strip():
                        full_text_parts.append(f"--- Page {page_num} ---\n{text}\n")
                
                doc.full_text = "\n".join(full_text_parts)
                
                # Chunk the document
                doc.chunks = self._chunk_text(
                    doc.full_text,
                    preserve_structure=True
                )
                
                # Store additional metadata
                doc.metadata = {
                    'num_pages': len(pdf_reader.pages),
                    'pdf_info': pdf_reader.metadata or {}
                }
        
        except Exception as e:
            raise ValueError(f"Error parsing PDF: {str(e)}")
        
        return doc
    
    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """Parse DOCX file."""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")
        
        doc = ParsedDocument(
            filename=file_path.name,
            file_type="DOCX"
        )
        
        try:
            docx_file = DocxDocument(file_path)
            
            # Extract metadata
            core_props = docx_file.core_properties
            doc.title = core_props.title or ""
            doc.author = core_props.author or ""
            if core_props.created:
                doc.date = core_props.created
            
            # Extract text with structure
            paragraphs = []
            current_section = None
            
            for para in docx_file.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                
                # Detect headings (section titles)
                if para.style.name.startswith('Heading'):
                    current_section = text
                    paragraphs.append(f"\n## {text}\n")
                else:
                    paragraphs.append(text)
            
            doc.full_text = "\n".join(paragraphs)
            
            # Chunk with structure preservation
            doc.chunks = self._chunk_text(
                doc.full_text,
                preserve_structure=True,
                sections=[p for p in docx_file.paragraphs if p.style.name.startswith('Heading')]
            )
            
            # Store additional metadata
            doc.metadata = {
                'num_paragraphs': len(docx_file.paragraphs),
                'styles_used': list(set(p.style.name for p in docx_file.paragraphs))
            }
        
        except Exception as e:
            raise ValueError(f"Error parsing DOCX: {str(e)}")
        
        return doc
    
    def _parse_txt(self, file_path: Path) -> ParsedDocument:
        """Parse TXT file."""
        doc = ParsedDocument(
            filename=file_path.name,
            file_type="TXT"
        )
        
        try:
            # Try different encodings
            encodings = ['utf-8', 'latin-1', 'cp1252']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as file:
                        doc.full_text = file.read()
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError(f"Could not decode file with any encoding: {encodings}")
            
            # Extract title from filename or first line
            if not doc.title:
                doc.title = file_path.stem
                first_lines = doc.full_text.split('\n')[:3]
                for line in first_lines:
                    if len(line.strip()) > 5 and len(line.strip()) < 100:
                        doc.title = line.strip()
                        break
            
            # Chunk the document
            doc.chunks = self._chunk_text(
                doc.full_text,
                preserve_structure=True
            )
            
            doc.metadata = {
                'encoding': encoding,
                'num_lines': len(doc.full_text.split('\n'))
            }
        
        except Exception as e:
            raise ValueError(f"Error parsing TXT: {str(e)}")
        
        return doc
    
    def _chunk_text(
        self,
        text: str,
        preserve_structure: bool = True,
        sections: Optional[List] = None
    ) -> List[DocumentChunk]:
        """
        Chunk text intelligently, preserving semantic boundaries when possible.
        
        Args:
            text: Full text to chunk
            preserve_structure: Attempt to chunk at paragraph/section boundaries
            sections: Optional list of section markers
        
        Returns:
            List of DocumentChunk objects
        """
        chunks = []
        
        if self.semantic_chunking and preserve_structure:
            # Try to chunk at natural boundaries (paragraphs, sections)
            chunks = self._semantic_chunk(text, sections)
        else:
            # Simple sliding window chunking
            chunks = self._sliding_window_chunk(text)
        
        return chunks
    
    def _semantic_chunk(
        self,
        text: str,
        sections: Optional[List] = None
    ) -> List[DocumentChunk]:
        """Chunk text at semantic boundaries (paragraphs, sections)."""
        chunks = []
        
        # Split by double newlines (paragraphs) first
        paragraphs = re.split(r'\n\s*\n+', text)
        
        current_chunk = ""
        current_chunk_id = 0
        current_section = None
        char_offset = 0
        
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            
            # Detect section headers
            section_match = re.match(r'^##?\s*(.+)$', para)
            if section_match:
                current_section = section_match.group(1)
                # If we have accumulated content, save it as a chunk
                if current_chunk:
                    chunks.append(DocumentChunk(
                        content=current_chunk.strip(),
                        chunk_id=current_chunk_id,
                        section=current_section,
                        start_char=char_offset - len(current_chunk),
                        end_char=char_offset
                    ))
                    current_chunk_id += 1
                    current_chunk = ""
                continue
            
            # If adding this paragraph would exceed chunk size, save current chunk
            if current_chunk and len(current_chunk) + len(para) > self.chunk_size:
                chunks.append(DocumentChunk(
                    content=current_chunk.strip(),
                    chunk_id=current_chunk_id,
                    section=current_section,
                    start_char=char_offset - len(current_chunk),
                    end_char=char_offset
                ))
                current_chunk_id += 1
                
                # Start new chunk with overlap if needed
                if self.chunk_overlap > 0 and current_chunk:
                    overlap_text = current_chunk[-self.chunk_overlap:]
                    current_chunk = overlap_text + "\n\n" + para
                else:
                    current_chunk = para
            else:
                if current_chunk:
                    current_chunk += "\n\n" + para
                else:
                    current_chunk = para
            
            char_offset += len(para) + 2  # +2 for \n\n
        
        # Add remaining chunk
        if current_chunk.strip():
            chunks.append(DocumentChunk(
                content=current_chunk.strip(),
                chunk_id=current_chunk_id,
                section=current_section,
                start_char=char_offset - len(current_chunk),
                end_char=char_offset
            ))
        
        return chunks
    
    def _sliding_window_chunk(self, text: str) -> List[DocumentChunk]:
        """Simple sliding window chunking with overlap."""
        chunks = []
        text_length = len(text)
        start = 0
        chunk_id = 0
        
        while start < text_length:
            end = min(start + self.chunk_size, text_length)
            chunk_text = text[start:end]
            
            # Try to end at sentence boundary
            if end < text_length:
                sentence_end = max(
                    chunk_text.rfind('. '),
                    chunk_text.rfind('.\n'),
                    chunk_text.rfind('\n\n')
                )
                if sentence_end > self.chunk_size * 0.5:  # At least 50% of chunk size
                    chunk_text = chunk_text[:sentence_end + 1]
                    end = start + sentence_end + 1
            
            chunks.append(DocumentChunk(
                content=chunk_text.strip(),
                chunk_id=chunk_id,
                start_char=start,
                end_char=end
            ))
            
            chunk_id += 1
            start = end - self.chunk_overlap
        
        return chunks
    
    def _parse_pdf_date(self, date_str: str) -> Optional[datetime]:
        """Parse PDF date string format (D:YYYYMMDDHHmmSSOHH'mm')."""
        try:
            # PDF date format: D:YYYYMMDDHHmmSSOHH'mm'
            date_str = date_str.replace('D:', '')
            if len(date_str) >= 14:
                year = int(date_str[0:4])
                month = int(date_str[4:6])
                day = int(date_str[6:8])
                hour = int(date_str[8:10]) if len(date_str) > 8 else 0
                minute = int(date_str[10:12]) if len(date_str) > 10 else 0
                second = int(date_str[12:14]) if len(date_str) > 12 else 0
                return datetime(year, month, day, hour, minute, second)
        except (ValueError, IndexError):
            pass
        return None


def parse_document(file_path: str, **kwargs) -> ParsedDocument:
    """
    Convenience function to parse a document.
    
    Args:
        file_path: Path to the document file
        **kwargs: Additional arguments to pass to DocumentParser
    
    Returns:
        ParsedDocument object
    """
    parser = DocumentParser(**kwargs)
    return parser.parse_file(file_path)
